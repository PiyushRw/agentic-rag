import hashlib
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class DocumentLoadError(Exception):
    """Raised when a file cannot be turned into text."""


@dataclass
class PageText:
    page_number: int | None  # 1-based for PDFs, None for DOCX/plain text
    text: str


def compute_file_hash(path: Path) -> str:
    sha256 = hashlib.sha256()
    with Path(path).open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            sha256.update(block)
    return sha256.hexdigest()


def _load_pdf(path: Path) -> list[PageText]:
    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise DocumentLoadError("PDF is password-protected.")
        pages = []
        for number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(PageText(page_number=number, text=text))
        return pages
    except DocumentLoadError:
        raise
    except Exception as exc:
        raise DocumentLoadError(f"Could not read PDF: {exc}") from exc


def _load_docx(path: Path) -> list[PageText]:
    try:
        doc = DocxDocument(str(path))
    except Exception as exc:
        raise DocumentLoadError(f"Could not read DOCX: {exc}") from exc

    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n\n".join(parts)
    if not text.strip():
        return []
    return [PageText(page_number=None, text=text)]


def _load_text(path: Path) -> list[PageText]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    if not text.strip():
        return []
    return [PageText(page_number=None, text=text)]


_LOADERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".txt": _load_text,
    ".md": _load_text,
}


def load_document(path: Path) -> list[PageText]:
    path = Path(path)
    if not path.is_file():
        raise DocumentLoadError(f"File not found: {path.name}")

    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentLoadError(f"Unsupported file type '{extension}'. Supported: {supported}")

    pages = _LOADERS[extension](path)
    if not pages:
        raise DocumentLoadError(
            f"No extractable text in '{path.name}'. "
            "If this is a scanned PDF, it needs OCR, which this app doesn't do."
        )
    return pages